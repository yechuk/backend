#!/usr/bin/env python
"""
Extract the 2022 FA warning-factors workbook into a structured CSV.

Default behavior:
- reads only the "전체_타자+투수" sheet
- treats row 3 as the header row
- keeps player rows until the first fully empty row
- converts selected descriptive fields into structured columns
"""
from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "data" / "FA_2022_경고요인_v2.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "fa_2022_warning_factors.csv"
DEFAULT_SHEET = "전체_타자+투수"

SOURCE_HEADERS = [
    "선수",
    "구분",
    "포지션",
    "이전팀",
    "계약팀",
    "나이",
    "WAR (2021)",
    "실제AAV ($M)",
    "포지션 희소성",
    "Boras 여부",
    "나이 신호",
]

OUTPUT_HEADERS = [
    "player_name",
    "player_type",
    "position",
    "previous_team",
    "contract_team",
    "age",
    "war_2021",
    "actual_aav_musd",
    "position_scarcity_level",
    "position_scarcity_war_threshold",
    "position_scarcity_fa_count",
    "position_scarcity_has_war_data",
    "is_boras",
    "age_signal_level",
    "age_signal_reference_band",
    "age_signal_is_aging_risk",
]


def clean_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return re.sub(r"\s+", " ", value).strip()


def parse_is_boras(value: Any) -> int:
    cleaned = clean_text(value)
    if cleaned == "🟢 Non-Boras":
        return 0
    if cleaned == "🟡 Boras":
        return 1
    raise ValueError(f"Unexpected Boras value: {value!r}")


def parse_position_scarcity(value: Any) -> dict[str, Any]:
    cleaned = clean_text(value)
    if not isinstance(cleaned, str) or not cleaned:
        raise ValueError(f"Unexpected position scarcity value: {value!r}")

    if "과열 경고" in cleaned:
        level = "overheated"
    elif "주의" in cleaned:
        level = "caution"
    elif "안정" in cleaned:
        level = "stable"
    else:
        raise ValueError(f"Could not parse position scarcity level: {value!r}")

    if "WAR 데이터 없음" in cleaned:
        return {
            "position_scarcity_level": level,
            "position_scarcity_war_threshold": None,
            "position_scarcity_fa_count": None,
            "position_scarcity_has_war_data": 0,
        }

    threshold_match = re.search(r"WAR\s+([0-9.]+)\+\s*기준", cleaned)
    count_match = re.search(r"FA\s+([0-9]+)명", cleaned)
    if threshold_match is None or count_match is None:
        raise ValueError(f"Could not parse position scarcity details: {value!r}")

    return {
        "position_scarcity_level": level,
        "position_scarcity_war_threshold": float(threshold_match.group(1)),
        "position_scarcity_fa_count": int(count_match.group(1)),
        "position_scarcity_has_war_data": 1,
    }


def parse_age_signal(value: Any) -> dict[str, Any]:
    cleaned = clean_text(value)
    if not isinstance(cleaned, str) or not cleaned:
        raise ValueError(f"Unexpected age signal value: {value!r}")

    if "전성기 프리미엄" in cleaned:
        level = "prime_premium"
    elif "에이징 리스크" in cleaned:
        level = "aging_risk"
    elif "해당없음" in cleaned:
        level = "none"
    else:
        raise ValueError(f"Could not parse age signal level: {value!r}")

    if "피크" in cleaned:
        reference_band = "prime_peak"
    elif "안정 구간" in cleaned:
        reference_band = "stable_30_34"
    elif "급락 구간" in cleaned:
        reference_band = "aging_decline"
    else:
        raise ValueError(f"Could not parse age signal reference band: {value!r}")

    return {
        "age_signal_level": level,
        "age_signal_reference_band": reference_band,
        "age_signal_is_aging_risk": 1 if level == "aging_risk" else 0,
    }


def validate_headers(row: list[Any]) -> None:
    cleaned_headers = [clean_text(cell) for cell in row]
    if cleaned_headers != SOURCE_HEADERS:
        raise ValueError(
            "Unexpected source headers.\n"
            f"Expected: {SOURCE_HEADERS}\n"
            f"Actual:   {cleaned_headers}"
        )


def iter_source_rows(input_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]

    header_row = next(
        worksheet.iter_rows(min_row=3, max_row=3, min_col=1, max_col=11, values_only=True)
    )
    validate_headers(list(header_row))

    records: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=4, min_col=1, max_col=11, values_only=True):
        values = list(row)
        if all(value is None or str(value).strip() == "" for value in values):
            break

        scarcity = parse_position_scarcity(values[8])
        age_signal = parse_age_signal(values[10])
        records.append(
            {
                "player_name": clean_text(values[0]),
                "player_type": clean_text(values[1]),
                "position": clean_text(values[2]),
                "previous_team": clean_text(values[3]),
                "contract_team": clean_text(values[4]),
                "age": values[5],
                "war_2021": values[6],
                "actual_aav_musd": values[7],
                **scarcity,
                "is_boras": parse_is_boras(values[9]),
                **age_signal,
            }
        )

    return records


def write_csv(output_path: Path, rows: list[dict[str, Any]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract the 2022 FA warning-factors workbook into a structured CSV."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Input workbook path.")
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help="Workbook sheet to extract. Defaults to 전체_타자+투수.",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output CSV path.")
    return parser


def main() -> None:
    parser = build_argument_parser()
    args = parser.parse_args()

    rows = iter_source_rows(args.input, args.sheet)
    write_csv(args.output, rows)
    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
