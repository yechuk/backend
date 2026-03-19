#!/usr/bin/env python
"""
Consolidate MLB free agent yearly sheets into one clean workbook.

Usage:
    python scripts/consolidate_mlb_free_agency.py
    python scripts/consolidate_mlb_free_agency.py --input path/to/source.xlsx --output path/to/output.xlsx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Iterator

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "data" / "MLB-Free Agency 1991-2026.xls.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "MLB-Free Agency 1991-2026_consolidated.xlsx"
OUTPUT_SHEET_NAME = "free_agents_1991_2026"
YEAR_RANGE = range(1991, 2027)
SECTION_MARKERS = {
    "international": "international",
    "retired": "retired",
    "retired / did not play": "retired",
}
INTERNATIONAL_CLUBS = {"NPB", "KBO", "MEX"}
RETIRED_NEW_CLUB_VALUES = {"dnp", "retired"}
NAME_SUFFIXES = {"Jr.", "Sr.", "II", "III", "IV"}

HEADER_MAP = {
    "Player": "Player",
    "Pos'n": "Position",
    "Age": "Age",
    "Qual Offer": "Qualifying Offer",
    "Old Club": "Old Club",
    "New Club": "New Club",
    "Years": "Years",
    "Guarantee": "Guarantee",
    "Term": "Term",
    "Option": "Option",
    "Opt Out": "Opt Out",
    "AAV": "AAV",
    "Player Agent": "Player Agent",
    "Club Owner": "New Club Owner",
    "Baseball Ops head / club GM": "New Club Baseball Ops Head / Club GM",
    "Details": "Details",
}

OUTPUT_COLUMNS = [
    "Year",
    "Status",
    "Player",
    "Position",
    "Age",
    "Qualifying Offer",
    "Old Club",
    "New Club",
    "Years",
    "Guarantee",
    "Term",
    "Option",
    "Opt Out",
    "AAV",
    "Player Agent",
    "New Club Owner",
    "New Club Baseball Ops Head / Club GM",
    "Details",
]


def clean_spacing(value: object) -> object:
    if not isinstance(value, str):
        return value
    return re.sub(r"\s+", " ", value).strip()


def normalize_header(value: object) -> str:
    cleaned = clean_spacing(value)
    if not isinstance(cleaned, str):
        raise ValueError(f"Unexpected header value: {value!r}")
    if cleaned.startswith("Age 7/1/"):
        return "Age"
    mapped = HEADER_MAP.get(cleaned)
    if mapped is None:
        raise ValueError(f"Unexpected source header: {cleaned!r}")
    return mapped


def normalize_player_name(value: object) -> object:
    cleaned = clean_spacing(value)
    if not isinstance(cleaned, str) or "," not in cleaned:
        return cleaned

    last_part, first_part = (clean_spacing(part) for part in cleaned.split(",", 1))
    if not isinstance(last_part, str) or not isinstance(first_part, str):
        return cleaned

    first_tokens = first_part.split()
    suffix = None
    if first_tokens and first_tokens[-1] in NAME_SUFFIXES:
        suffix = first_tokens.pop()
        first_part = " ".join(first_tokens)

    name_parts = [part for part in (first_part, last_part, suffix) if part]
    return " ".join(name_parts)


def is_remaining_free_agent_label(value: object) -> bool:
    cleaned = clean_spacing(value)
    return isinstance(cleaned, str) and cleaned.lower().startswith("remaining free agents")


def is_section_label(value: object) -> bool:
    cleaned = clean_spacing(value)
    if not isinstance(cleaned, str):
        return False
    lowered = cleaned.lower()
    return lowered in SECTION_MARKERS or lowered == "international" or is_remaining_free_agent_label(cleaned)


def iter_player_rows(sheet) -> Iterator[list[object]]:
    header_row_index = None
    header_names: list[str] = []
    current_section = "signed"

    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        first_16 = list(row[:16])
        if header_row_index is None:
            if first_16 and first_16[0] == "Player":
                header_row_index = row_index
                header_names = [normalize_header(value) for value in first_16]
            continue

        if not any(value is not None and str(value).strip() != "" for value in first_16):
            continue

        first_cell = clean_spacing(first_16[0])
        if is_section_label(first_cell):
            lowered = first_cell.lower() if isinstance(first_cell, str) else ""
            if lowered == "international":
                current_section = "international"
            elif lowered in SECTION_MARKERS:
                current_section = SECTION_MARKERS[lowered]
            elif is_remaining_free_agent_label(first_cell):
                current_section = "remaining_free_agent"
            continue

        if not isinstance(first_cell, str) or first_cell == "":
            continue

        record = {}
        for header_name, value in zip(header_names, first_16):
            cleaned_value = clean_spacing(value)
            if header_name == "Player":
                cleaned_value = normalize_player_name(cleaned_value)
            record[header_name] = cleaned_value

        status = derive_status(record.get("New Club"), current_section)
        yield [
            int(sheet.title),
            status,
            record.get("Player"),
            record.get("Position"),
            record.get("Age"),
            record.get("Qualifying Offer"),
            record.get("Old Club"),
            record.get("New Club"),
            record.get("Years"),
            record.get("Guarantee"),
            record.get("Term"),
            record.get("Option"),
            record.get("Opt Out"),
            record.get("AAV"),
            record.get("Player Agent"),
            record.get("New Club Owner"),
            record.get("New Club Baseball Ops Head / Club GM"),
            record.get("Details"),
        ]

    if header_row_index is None:
        raise ValueError(f"Could not find player header row in sheet {sheet.title!r}")


def derive_status(new_club: object, current_section: str) -> str:
    if current_section != "signed":
        return current_section

    if new_club in INTERNATIONAL_CLUBS:
        return "international"
    if new_club in RETIRED_NEW_CLUB_VALUES:
        return "retired"
    if new_club in (None, ""):
        return "remaining_free_agent"
    return "signed"


def build_workbook(input_path: Path, output_path: Path) -> tuple[int, dict[str, int]]:
    source_wb = load_workbook(input_path, data_only=True)
    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = OUTPUT_SHEET_NAME
    output_ws.append(OUTPUT_COLUMNS)
    output_ws.freeze_panes = "A2"
    output_ws.auto_filter.ref = f"A1:R1"

    row_count = 0
    status_counts = {
        "signed": 0,
        "international": 0,
        "retired": 0,
        "remaining_free_agent": 0,
    }

    for year in YEAR_RANGE:
        sheet = source_wb[str(year)]
        for record in iter_player_rows(sheet):
            output_ws.append(record)
            row_count += 1
            status_counts[record[1]] += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)
    return row_count, status_counts


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Consolidate MLB free agent sheets into one workbook.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Source workbook path.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output workbook path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    row_count, status_counts = build_workbook(args.input, args.output)
    print(f"Created: {args.output}")
    print(f"Rows: {row_count}")
    for status, count in status_counts.items():
        print(f"{status}: {count}")


if __name__ == "__main__":
    main()
