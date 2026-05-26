from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from .models import (
    MLBApiAavPrediction,
    MLBApiPerformanceValuePrediction,
    MLBApiRecommendedSimilarPlayer,
    MLBApiSimilarPlayer,
    MLBApiStatLine,
    MLBApiTeamDollarPerWar,
    MLBRosterEntry,
    MLBRosterPhoto,
)


class TeamDollarPerWarApiTests(TestCase):
    def setUp(self):
        MLBApiTeamDollarPerWar.objects.create(
            team='WSN',
            avg_payroll_m_3yr=117.0,
            avg_batting_war_3yr=5.966111490666666,
            avg_pitching_war_3yr=2.970400869666667,
            avg_team_war_3yr=8.936512360333333,
            n_years=3,
            avg_team_war_3yr_safe=8.936512360333333,
            dollar_per_war_millions=13.092355863494367,
            dollar_per_war=13092355.863494366,
        )

    def test_endpoint_returns_team_dollar_per_war_rows_from_db(self):
        response = self.client.get('/api/team-dollar-per-war/')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['source'], 'team_dollar_per_war.csv')
        self.assertEqual(payload['count'], 1)

        first_team = payload['teams'][0]
        self.assertEqual(first_team['team'], 'WSN')
        self.assertIsInstance(first_team['avg_payroll_m_3yr'], float)
        self.assertIsInstance(first_team['avg_batting_war_3yr'], float)
        self.assertIsInstance(first_team['avg_pitching_war_3yr'], float)
        self.assertIsInstance(first_team['avg_team_war_3yr'], float)
        self.assertIsInstance(first_team['n_years'], int)
        self.assertIsInstance(first_team['avg_team_war_3yr_safe'], float)
        self.assertIsInstance(first_team['dollar_per_war_millions'], float)
        self.assertIsInstance(first_team['dollar_per_war'], float)

    def test_endpoint_no_slash_alias_returns_same_payload(self):
        response = self.client.get('/api/team-dollar-per-war')
        slash_response = self.client.get('/api/team-dollar-per-war/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), slash_response.json())

    def test_endpoint_returns_empty_list_when_db_is_empty(self):
        MLBApiTeamDollarPerWar.objects.all().delete()

        response = self.client.get('/api/team-dollar-per-war/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {
            'source': 'team_dollar_per_war.csv',
            'count': 0,
            'teams': [],
        })


class LoadTeamDollarPerWarCommandTests(TestCase):
    def _write_csv(self, base_dir, body):
        data_dir = Path(base_dir) / 'data'
        data_dir.mkdir(parents=True, exist_ok=True)
        (data_dir / 'team_dollar_per_war.csv').write_text(body, encoding='utf-8')

    def test_command_loads_csv_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            self._write_csv(
                temp_dir,
                'Team,avg_payroll_M_3yr,avg_batting_war_3yr,avg_pitching_war_3yr,avg_team_war_3yr,n_years,avg_team_war_3yr_safe,dollar_per_war_M,dollar_per_war\n'
                'WSN,117.0,5.96,2.97,8.93,3,8.93,13.09,13090000\n'
                'TEX,121.13,8.48,4.18,12.66,3,12.66,9.56,9560000\n',
            )

            call_command('load_team_dollar_per_war', '--replace', base_dir=temp_dir, verbosity=0)

        self.assertEqual(MLBApiTeamDollarPerWar.objects.count(), 2)
        wsn = MLBApiTeamDollarPerWar.objects.get(team='WSN')
        self.assertEqual(wsn.n_years, 3)
        self.assertEqual(wsn.avg_payroll_m_3yr, 117.0)
        self.assertEqual(wsn.dollar_per_war_millions, 13.09)

    def test_command_replace_deletes_existing_rows_before_load(self):
        MLBApiTeamDollarPerWar.objects.create(team='OLD', dollar_per_war_millions=1.0)
        with TemporaryDirectory() as temp_dir:
            self._write_csv(
                temp_dir,
                'Team,avg_payroll_M_3yr,avg_batting_war_3yr,avg_pitching_war_3yr,avg_team_war_3yr,n_years,avg_team_war_3yr_safe,dollar_per_war_M,dollar_per_war\n'
                'WSN,117.0,5.96,2.97,8.93,3,8.93,13.09,13090000\n',
            )

            call_command('load_team_dollar_per_war', '--replace', base_dir=temp_dir, verbosity=0)

        self.assertFalse(MLBApiTeamDollarPerWar.objects.filter(team='OLD').exists())
        self.assertEqual(MLBApiTeamDollarPerWar.objects.count(), 1)

    def test_command_updates_existing_team_without_duplicate(self):
        MLBApiTeamDollarPerWar.objects.create(team='WSN', dollar_per_war_millions=1.0)
        with TemporaryDirectory() as temp_dir:
            self._write_csv(
                temp_dir,
                'Team,avg_payroll_M_3yr,avg_batting_war_3yr,avg_pitching_war_3yr,avg_team_war_3yr,n_years,avg_team_war_3yr_safe,dollar_per_war_M,dollar_per_war\n'
                'WSN,117.0,5.96,2.97,8.93,3,8.93,13.09,13090000\n',
            )

            call_command('load_team_dollar_per_war', base_dir=temp_dir, verbosity=0)

        self.assertEqual(MLBApiTeamDollarPerWar.objects.count(), 1)
        self.assertEqual(MLBApiTeamDollarPerWar.objects.get(team='WSN').dollar_per_war_millions, 13.09)

    def test_command_errors_when_csv_is_missing(self):
        with TemporaryDirectory() as temp_dir:
            with self.assertRaises(CommandError):
                call_command('load_team_dollar_per_war', base_dir=temp_dir, verbosity=0)

    def test_command_errors_on_missing_team(self):
        with TemporaryDirectory() as temp_dir:
            self._write_csv(
                temp_dir,
                'Team,avg_payroll_M_3yr,avg_batting_war_3yr,avg_pitching_war_3yr,avg_team_war_3yr,n_years,avg_team_war_3yr_safe,dollar_per_war_M,dollar_per_war\n'
                ',117.0,5.96,2.97,8.93,3,8.93,13.09,13090000\n',
            )

            with self.assertRaises(CommandError):
                call_command('load_team_dollar_per_war', base_dir=temp_dir, verbosity=0)

    def test_command_errors_on_duplicate_team(self):
        with TemporaryDirectory() as temp_dir:
            self._write_csv(
                temp_dir,
                'Team,avg_payroll_M_3yr,avg_batting_war_3yr,avg_pitching_war_3yr,avg_team_war_3yr,n_years,avg_team_war_3yr_safe,dollar_per_war_M,dollar_per_war\n'
                'WSN,117.0,5.96,2.97,8.93,3,8.93,13.09,13090000\n'
                'WSN,118.0,5.96,2.97,8.93,3,8.93,14.09,14090000\n',
            )

            with self.assertRaises(CommandError):
                call_command('load_team_dollar_per_war', base_dir=temp_dir, verbosity=0)


class TeamApiTests(TestCase):
    def setUp(self):
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2023,
            team='BAL',
            player_name='Kyle Gibson',
            name_ascii='Kyle Gibson',
            external_player_id='10123',
            mlbam_id='502043',
            age=35,
            war=2.66114687919617,
            raw_stats={
                'Season': '2023',
                'Name': 'Kyle Gibson',
                'Team': 'BAL',
                'HR/9': '1.078125',
                'K%': '0.19454771',
                'BB': '55',
                'IP': '192.0',
                'FIP': '4.13004035949707',
                'GS': '33',
                'G': '33',
                'Age': '35',
                'WAR': '2.66114687919617',
                'NameASCII': 'Kyle Gibson',
                'PlayerId': '10123',
                'MLBAMID': '502043',
            },
        )
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2023,
            team='ATL',
            player_name='Max Fried',
            name_ascii='Max Fried',
            external_player_id='20456',
            mlbam_id='608331',
            age=29,
            war=3.5,
            raw_stats={
                'Season': '2023',
                'Name': 'Max Fried',
                'Team': 'ATL',
                'HR/9': '0.8',
                'K%': '0.2400',
                'BB': '40',
                'IP': '180.0',
                'FIP': '3.10',
                'GS': '30',
                'G': '30',
                'Age': '29',
                'WAR': '3.5',
                'NameASCII': 'Max Fried',
                'PlayerId': '20456',
                'MLBAMID': '608331',
            },
        )

    def test_team_list_endpoint_returns_2023_pitching_teams(self):
        response = self.client.get('/api/teams/?season=2023&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2023)
        self.assertEqual(payload['view'], 'pitching')
        self.assertTrue(any(team['code'] == 'BAL' for team in payload['teams']))

    def test_team_players_endpoint_returns_baltimore_pitchers(self):
        response = self.client.get('/api/teams/BAL/players/?season=2023&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['team'], 'BAL')
        self.assertTrue(any(player['player_id'] == '502043' for player in payload['players']))

    def test_team_player_detail_endpoint_returns_kyle_gibson(self):
        response = self.client.get('/api/teams/BAL/players/502043/?season=2023&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['player']['name'], 'Kyle Gibson')
        self.assertEqual(payload['player']['team'], 'BAL')
        self.assertEqual(payload['player']['player_id'], '502043')
        self.assertEqual(payload['player']['stats']['games_started'], 33)

    def test_team_players_endpoint_rounds_batting_wrc_plus(self):
        MLBApiStatLine.objects.create(
            stat_view='batting',
            season=2023,
            team='BAL',
            player_name='Adley Rutschman',
            name_ascii='Adley Rutschman',
            external_player_id='20001',
            mlbam_id='668939',
            age=25,
            war=5.4,
            raw_stats={
                'Season': '2023',
                'Name': 'Adley Rutschman',
                'Team': 'BAL',
                'G': '154',
                'PA': '687',
                'HR': '20',
                'ISO': '0.192',
                'BB%': '0.133',
                'K%': '0.178',
                'wOBA': '0.374',
                'wRC+': '127.6',
                'SB': '1',
                'BABIP': '0.317',
                'Age': '25',
                'WAR': '5.4',
                'NameASCII': 'Adley Rutschman',
                'PlayerId': '20001',
                'MLBAMID': '668939',
            },
        )

        response = self.client.get('/api/teams/BAL/players/?season=2023&view=batting')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        adley = next(player for player in payload['players'] if player['player_id'] == '668939')
        self.assertEqual(adley['stats']['wrc_plus'], 128)

    def test_team_player_detail_endpoint_supports_enriched_pitching_fields(self):
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2022,
            team='ATL',
            player_name='A.J. Minter',
            name_ascii='A.J. Minter',
            external_player_id='621345',
            mlbam_id='621345',
            age=28,
            war=2.0,
            raw_stats={
                'Season': '2022',
                'Name': 'A.J. Minter',
                'Team': 'ATL',
                'Position': 'P',
                'Height': '6\' 0"',
                'Weight': '215',
                'Bats': 'L',
                'Throws': 'L',
                'ERA': '2.06',
                'FIP': '2.53',
                'WHIP': '0.96',
                'K/9': '10.67',
                'BB/9': '3.09',
                'IP': '70.0',
                'SO': '83',
                'xERA': '2.65',
                'xFIP': '2.48',
                'LOB%': '81.2',
                'BABIP': '0.245',
                'HR/9': '0.64',
                'velocity': '96.4',
                'Age': '28',
                'WAR': '2.0',
                'PlayerId': '621345',
                'MLBAMID': '621345',
            },
        )

        response = self.client.get('/api/teams/ATL/players/621345/?season=2022&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['player']['position'], 'P')
        self.assertEqual(payload['player']['height'], '6\' 0"')
        self.assertEqual(payload['player']['weight'], 215)
        self.assertEqual(payload['player']['stats']['era'], 2.06)
        self.assertEqual(payload['player']['stats']['k_per_9'], 10.67)
        self.assertEqual(payload['player']['stats']['x_fip'], 2.48)

    def test_team_player_detail_endpoint_returns_five_year_history_and_photo(self):
        MLBRosterPhoto.objects.create(
            team_name='Atlanta Braves',
            player_name='A.J. Minter',
            normalized_player_name='ajminter',
            original_filename='A.J. Minter.jpeg',
            content_type='image/jpeg',
            image_data=b'fake-image-bytes',
            byte_size=16,
        )
        MLBRosterEntry.objects.create(
            season=2022,
            team_id=144,
            team_name='Atlanta Braves',
            team_abbreviation='ATL',
            league_name='National League',
            division_name='National League East',
            player_id=621345,
            player_name='A.J. Minter',
            player_link='/api/v1/people/621345',
            jersey_number='33',
            position_code='1',
            position_name='Pitcher',
            position_type='Pitcher',
            position_abbreviation='P',
            status_code='A',
            status_description='Active',
            raw_data={},
        )
        for season, era, war, team in [
            (2022, '2.06', '2.0', 'ATL'),
            (2021, '3.30', '1.2', 'ATL'),
            (2020, '0.83', '0.8', 'ATL'),
            (2019, '7.06', '-0.4', 'ATL'),
            (2018, '3.23', '1.4', 'ATL'),
        ]:
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=season,
                team=team,
                player_name='A.J. Minter',
                name_ascii='A.J. Minter',
                external_player_id='18655',
                mlbam_id='621345',
                age=28,
                war=float(war),
                raw_stats={
                    'Season': str(season),
                    'Name': 'A.J. Minter',
                    'Team': team,
                    'Position': 'P',
                    'Height': '6\' 0"',
                    'Weight': '215',
                    'Bats': 'L',
                    'Throws': 'L',
                    'ERA': era,
                    'WAR': war,
                    'PlayerId': '18655',
                    'MLBAMID': '621345',
                },
            )

        response = self.client.get('/api/teams/ATL/players/621345/?view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['history_count'], 5)
        self.assertEqual(
            [row['season'] for row in payload['history']],
            [2022, 2021, 2020, 2019, 2018],
        )
        self.assertEqual(
            payload['player']['photo_url'],
            '/api/rosters/ATL/players/621345/photo/?season=2022',
        )
        self.assertEqual(payload['history'][0]['height'], '6\' 0"')
        self.assertEqual(payload['history'][0]['bats'], 'L')

    def test_team_player_detail_endpoint_prefers_roster_season_aggregate_stat_line(self):
        MLBRosterEntry.objects.create(
            season=2022,
            team_id=144,
            team_name='Atlanta Braves',
            team_abbreviation='ATL',
            league_name='National League',
            division_name='National League East',
            player_id=445926,
            player_name='Jesse Chavez',
            player_link='/api/v1/people/445926',
            jersey_number='60',
            position_code='1',
            position_name='Pitcher',
            position_type='Pitcher',
            position_abbreviation='P',
            status_code='A',
            status_description='Active',
            raw_data={},
        )
        for season, team, games, war in [
            (2022, '- - -', '60', '0.72'),
            (2021, 'ATL', '30', '0.96'),
        ]:
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=season,
                team=team,
                player_name='Jesse Chavez',
                name_ascii='Jesse Chavez',
                external_player_id='5448',
                mlbam_id='445926',
                age=38,
                war=float(war),
                raw_stats={
                    'Season': str(season),
                    'Name': 'Jesse Chavez',
                    'Team': team,
                    'G': games,
                    'WAR': war,
                    'PlayerId': '5448',
                    'MLBAMID': '445926',
                },
            )

        response = self.client.get('/api/teams/ATL/players/445926/?view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2022)
        self.assertEqual(payload['player']['season'], 2022)
        self.assertEqual(payload['player']['team'], '- - -')
        self.assertEqual(payload['player']['stats']['games'], 60)
        self.assertEqual(payload['history_count'], 1)
        self.assertEqual([row['season'] for row in payload['history']], [2022])

        response = self.client.get('/api/teams/ATL/players/445926/?season=2022&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2022)
        self.assertEqual(payload['history_count'], 1)

    def test_team_player_detail_endpoint_with_season_returns_only_requested_season(self):
        for season, era, war, team in [
            (2022, '2.06', '2.0', 'ATL'),
            (2021, '3.30', '1.2', 'ATL'),
        ]:
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=season,
                team=team,
                player_name='A.J. Minter',
                name_ascii='A.J. Minter',
                external_player_id='18655',
                mlbam_id='621345',
                age=28,
                war=float(war),
                raw_stats={
                    'Season': str(season),
                    'Name': 'A.J. Minter',
                    'Team': team,
                    'Position': 'P',
                    'Height': '6\' 0"',
                    'Weight': '215',
                    'Bats': 'L',
                    'Throws': 'L',
                    'ERA': era,
                    'WAR': war,
                    'PlayerId': '18655',
                    'MLBAMID': '621345',
                },
            )

        response = self.client.get('/api/teams/ATL/players/621345/?season=2021&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2021)
        self.assertEqual(payload['player']['season'], 2021)
        self.assertEqual(payload['history_count'], 1)
        self.assertEqual(len(payload['history']), 1)
        self.assertEqual(payload['history'][0]['season'], 2021)

    def test_team_player_detail_endpoint_supports_enriched_batting_fields(self):
        MLBApiStatLine.objects.create(
            stat_view='batting',
            season=2022,
            team='ATL',
            player_name='Dansby Swanson',
            name_ascii='Dansby Swanson',
            external_player_id='621020',
            mlbam_id='621020',
            age=28,
            war=6.4,
            raw_stats={
                'Season': '2022',
                'Name': 'Dansby Swanson',
                'Team': 'ATL',
                'Position': 'SS',
                'Height': '6\' 1"',
                'Weight': '190',
                'Bats': 'R',
                'Throws': 'R',
                'AVG': '0.277',
                'OPS': '0.776',
                'HR': '25',
                'RBI': '96',
                'wRC+': '116',
                'wOBA': '0.330',
                'BABIP': '0.348',
                'ISO': '0.179',
                'BB%': '0.087',
                'K%': '0.261',
                'ExitVelocity': '89.5',
                'LaunchAngle': '14.2',
                'Age': '28',
                'WAR': '6.4',
                'PlayerId': '621020',
                'MLBAMID': '621020',
            },
        )

        response = self.client.get('/api/teams/ATL/players/621020/?season=2022&view=batting')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['player']['position'], 'SS')
        self.assertEqual(payload['player']['stats']['avg'], 0.277)
        self.assertEqual(payload['player']['stats']['ops'], 0.776)
        self.assertEqual(payload['player']['stats']['rbi'], 96)
        self.assertEqual(payload['player']['stats']['exit_velocity'], 89.5)

    def test_team_player_detail_endpoint_includes_similar_players(self):
        MLBApiSimilarPlayer.objects.create(
            stat_view='pitching',
            source_player_name='A.J. Minter',
            source_name_ascii='ajminter',
            source_mlbam_id='621345',
            source_external_player_id='18655',
            similar_player_name='Matthew Boyd',
            similar_name_ascii='matthewboyd',
            similar_mlbam_id='571510',
            similar_external_player_id='15440',
            similar_team='DET',
            similarity_score=48,
            rank=1,
        )
        MLBApiRecommendedSimilarPlayer.objects.create(
            stat_view='pitching',
            source_player_name='A.J. Minter',
            source_name_ascii='ajminter',
            source_mlbam_id='621345',
            source_external_player_id='18655',
            source_player_position='P',
            source_player_age=28,
            similar_player_name='Joely Rodríguez',
            similar_name_ascii='joelyrodriguez',
            similar_mlbam_id='570257',
            similar_external_player_id='11487',
            similar_team='NYM',
            similar_player_position='P',
            similar_player_age=30,
            similarity_score='0.912345678',
            rank=1,
        )
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2022,
            team='ATL',
            player_name='A.J. Minter',
            name_ascii='A.J. Minter',
            external_player_id='18655',
            mlbam_id='621345',
            age=28,
            war=2.0,
            raw_stats={
                'Season': '2022',
                'Name': 'A.J. Minter',
                'Team': 'ATL',
                'PlayerId': '18655',
                'MLBAMID': '621345',
            },
        )

        response = self.client.get('/api/teams/ATL/players/621345/?season=2022&view=pitching')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload['euclidean_similar_players']), 1)
        self.assertEqual(payload['euclidean_similar_players'][0]['player_name'], 'Matthew Boyd')
        self.assertEqual(payload['euclidean_similar_players'][0]['similarity_score'], 48)
        self.assertEqual(
            payload['euclidean_similar_players'][0]['teams_detail_url'],
            '/api/teams/DET/players/571510/?view=pitching',
        )
        self.assertEqual(len(payload['tabnet_similar_players']), 1)
        self.assertEqual(payload['tabnet_similar_players'][0]['player_name'], 'Joely Rodríguez')
        self.assertAlmostEqual(payload['tabnet_similar_players'][0]['similarity_score'], 0.912345678)
        self.assertEqual(payload['tabnet_similar_players'][0]['position'], 'P')
        self.assertEqual(payload['tabnet_similar_players'][0]['age'], 30)
        self.assertEqual(payload['similar_players'], payload['euclidean_similar_players'])
        self.assertEqual(payload['similar_player_recommendations'], payload['tabnet_similar_players'])
        self.assertIsNone(payload['predicted_aav'])
        self.assertEqual(payload['predicted_aav_meta']['source'], 'M1')
        self.assertFalse(payload['predicted_aav_meta']['available'])

    def test_team_player_detail_endpoint_includes_predicted_aav_for_matching_batter(self):
        MLBApiAavPrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M1',
            source_file='M1_2022_Predictions.xlsx',
            player_name='Dansby Swanson',
            name_ascii='dansbyswanson',
            team_code_raw='ATL',
            position_raw='SS',
            actual_aav_millions='21.25',
            predicted_aav_millions='19.75',
            prediction_error_millions='-1.50',
        )
        MLBApiPerformanceValuePrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M2',
            source_file='M2_WAR_Value_2022_Full.xlsx',
            player_name='Dansby Swanson',
            name_ascii='dansbyswanson',
            player_type='hitter',
            current_team='ATL',
            target_team='ATL',
            war_2022='6.400',
            predicted_war_avg='3.125',
            actual_war_avg='2.500',
            dollars_per_war_millions='7.10',
            predicted_value_millions='22.19',
        )
        MLBApiStatLine.objects.create(
            stat_view='batting',
            season=2022,
            team='ATL',
            player_name='Dansby Swanson',
            name_ascii='Dansby Swanson',
            external_player_id='30002',
            mlbam_id='621020',
            age=28,
            war=6.4,
            raw_stats={
                'Season': '2022',
                'Name': 'Dansby Swanson',
                'Team': 'ATL',
                'G': '162',
                'PA': '696',
                'HR': '25',
                'RBI': '96',
                'AVG': '0.277',
                'OPS': '0.776',
                'ISO': '0.179',
                'BB%': '0.087',
                'K%': '0.261',
                'wOBA': '0.330',
                'wRC+': '116',
                'SB': '18',
                'BABIP': '0.348',
                'Age': '28',
                'WAR': '6.4',
                'NameASCII': 'Dansby Swanson',
                'PlayerId': '30002',
                'MLBAMID': '621020',
            },
        )

        response = self.client.get('/api/teams/ATL/players/621020/?season=2022&view=batting')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['predicted_aav'], 19.75)
        self.assertEqual(payload['predicted_market_value'], 19.75)
        self.assertEqual(
            payload['predicted_market_value_meta'],
            {
                'source': 'M1',
                'source_file': 'M1_2022_Predictions.xlsx',
                'season': 2022,
                'view': 'batting',
                'basis': 'market_aav',
                'unit': 'USD_M',
                'available': True,
            },
        )
        self.assertEqual(payload['predicted_performance_value'], 22.19)
        self.assertEqual(
            payload['predicted_performance_value_meta'],
            {
                'source': 'M2',
                'source_file': 'M2_WAR_Value_2022_Full.xlsx',
                'season': 2022,
                'view': 'batting',
                'basis': 'war_value_aav',
                'target_team': 'ATL',
                'unit': 'USD_M',
                'available': True,
            },
        )
        self.assertEqual(
            payload['predicted_aav_meta'],
            {
                'source': 'M1',
                'source_file': 'M1_2022_Predictions.xlsx',
                'season': 2022,
                'view': 'batting',
                'unit': 'USD_M',
                'available': True,
            },
        )

    def test_team_player_detail_endpoint_rejects_external_player_id_lookup(self):
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2022,
            team='ATL',
            player_name='A.J. Minter',
            name_ascii='A.J. Minter',
            external_player_id='18655',
            mlbam_id='621345',
            age=28,
            war=2.0,
            raw_stats={
                'Season': '2022',
                'Name': 'A.J. Minter',
                'Team': 'ATL',
                'PlayerId': '18655',
                'MLBAMID': '621345',
            },
        )

        response = self.client.get('/api/teams/ATL/players/18655/?season=2022&view=pitching')

        self.assertEqual(response.status_code, 404)


class RosterApiTests(TestCase):
    def setUp(self):
        MLBRosterPhoto.objects.create(
            team_name='Atlanta Braves',
            player_name='A.J. Minter',
            normalized_player_name='ajminter',
            original_filename='A.J. Minter.jpeg',
            content_type='image/jpeg',
            image_data=b'fake-image-bytes',
            byte_size=16,
        )
        MLBApiStatLine.objects.create(
            stat_view='pitching',
            season=2022,
            team='ATL',
            player_name='A.J. Minter',
            name_ascii='A.J. Minter',
            external_player_id='30001',
            mlbam_id='621345',
            age=28,
            war=1.8,
            raw_stats={
                'Season': '2022',
                'Name': 'A.J. Minter',
                'Team': 'ATL',
                'HR/9': '0.70',
                'K%': '0.313',
                'BB': '24',
                'IP': '70.0',
                'FIP': '2.53',
                'GS': '0',
                'G': '61',
                'Age': '28',
                'WAR': '1.8',
                'NameASCII': 'A.J. Minter',
                'PlayerId': '30001',
                'MLBAMID': '621345',
            },
        )
        MLBApiStatLine.objects.create(
            stat_view='batting',
            season=2022,
            team='ATL',
            player_name='Dansby Swanson',
            name_ascii='Dansby Swanson',
            external_player_id='30002',
            mlbam_id='621020',
            age=28,
            war=6.4,
            raw_stats={
                'Season': '2022',
                'Name': 'Dansby Swanson',
                'Team': 'ATL',
                'G': '162',
                'PA': '696',
                'HR': '25',
                'ISO': '0.179',
                'BB%': '0.087',
                'K%': '0.261',
                'wOBA': '0.330',
                'wRC+': '116',
                'SB': '18',
                'BABIP': '0.348',
                'Age': '28',
                'WAR': '6.4',
                'NameASCII': 'Dansby Swanson',
                'PlayerId': '30002',
                'MLBAMID': '621020',
            },
        )
        MLBRosterEntry.objects.create(
            season=2022,
            team_id=144,
            team_name='Atlanta Braves',
            team_abbreviation='ATL',
            league_name='National League',
            division_name='National League East',
            player_id=621345,
            player_name='A.J. Minter',
            player_link='/api/v1/people/621345',
            jersey_number='33',
            position_code='1',
            position_name='Pitcher',
            position_type='Pitcher',
            position_abbreviation='P',
            status_code='A',
            status_description='Active',
            raw_data={},
        )
        MLBRosterEntry.objects.create(
            season=2022,
            team_id=144,
            team_name='Atlanta Braves',
            team_abbreviation='ATL',
            league_name='National League',
            division_name='National League East',
            player_id=621020,
            player_name='Dansby Swanson',
            player_link='/api/v1/people/621020',
            jersey_number='7',
            position_code='6',
            position_name='Shortstop',
            position_type='Infielder',
            position_abbreviation='SS',
            status_code='A',
            status_description='Active',
            raw_data={},
        )
        MLBRosterEntry.objects.create(
            season=2022,
            team_id=110,
            team_name='Baltimore Orioles',
            team_abbreviation='BAL',
            league_name='American League',
            division_name='American League East',
            player_id=660271,
            player_name='Adley Rutschman',
            player_link='/api/v1/people/660271',
            jersey_number='35',
            position_code='2',
            position_name='Catcher',
            position_type='Catcher',
            position_abbreviation='C',
            status_code='A',
            status_description='Active',
            raw_data={},
        )

    def test_roster_list_endpoint_returns_2022_teams(self):
        response = self.client.get('/api/rosters/?season=2022')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2022)
        self.assertEqual(payload['count'], 2)
        self.assertTrue(any(team['code'] == 'ATL' for team in payload['teams']))

    def test_team_roster_endpoint_returns_atlanta_roster(self):
        response = self.client.get('/api/rosters/ATL/?season=2022')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['team'], 'ATL')
        self.assertEqual(payload['count'], 2)
        self.assertTrue(any(player['player_name'] == 'A.J. Minter' for player in payload['players']))
        self.assertEqual(payload['stat_seasons'], {'batting': 2022, 'pitching': 2022})
        players_by_name = {player['player_name']: player for player in payload['players']}
        self.assertEqual(
            players_by_name['A.J. Minter']['photo_url'],
            '/api/rosters/ATL/players/621345/photo/?season=2022',
        )
        self.assertIsNone(players_by_name['Dansby Swanson']['photo_url'])
        self.assertEqual(players_by_name['A.J. Minter']['pitching']['mlbam_id'], '621345')
        self.assertEqual(players_by_name['A.J. Minter']['pitching']['stats']['games'], 61)
        self.assertIsNone(players_by_name['A.J. Minter']['batting'])
        self.assertEqual(players_by_name['Dansby Swanson']['batting']['mlbam_id'], '621020')
        self.assertEqual(players_by_name['Dansby Swanson']['batting']['stats']['home_runs'], 25)
        self.assertIsNone(players_by_name['Dansby Swanson']['pitching'])

    def test_roster_player_photo_endpoint_returns_image(self):
        response = self.client.get('/api/rosters/ATL/players/621345/photo/?season=2022')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'image/jpeg')
        self.assertEqual(response.content, b'fake-image-bytes')

    def test_roster_player_photo_endpoint_returns_404_when_missing(self):
        MLBRosterPhoto.objects.all().delete()
        response = self.client.get('/api/rosters/ATL/players/621345/photo/?season=2022')
        self.assertEqual(response.status_code, 404)

    def test_roster_player_detail_endpoint_returns_requested_season_stats(self):
        stat_line = MLBApiStatLine.objects.get(stat_view='pitching', season=2022, mlbam_id='621345')
        stat_line.external_player_id = '18655'
        stat_line.raw_stats['PlayerId'] = '18655'
        stat_line.save(update_fields=['external_player_id', 'raw_stats', 'updated_at'])

        response = self.client.get('/api/rosters/ATL/players/621345/?season=2022')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['season'], 2022)
        self.assertEqual(payload['roster_season'], 2022)
        self.assertEqual(payload['player']['player_id'], 621345)
        self.assertEqual(payload['player']['player_name'], 'A.J. Minter')
        self.assertEqual(payload['history_count'], 1)
        self.assertEqual(payload['history'][0]['season'], 2022)
        self.assertIsNone(payload['history'][0]['batting'])
        self.assertEqual(payload['history'][0]['pitching']['mlbam_id'], '621345')
        self.assertEqual(payload['history'][0]['pitching']['source_player_id'], '621345')
        self.assertEqual(payload['history'][0]['pitching']['source_external_player_id'], '18655')
        self.assertEqual(payload['history'][0]['pitching']['stats']['games'], 61)

    def test_roster_player_detail_endpoint_returns_five_year_history_without_season(self):
        stat_line = MLBApiStatLine.objects.get(stat_view='pitching', season=2022, mlbam_id='621345')
        stat_line.external_player_id = '18655'
        stat_line.raw_stats['PlayerId'] = '18655'
        stat_line.save(update_fields=['external_player_id', 'raw_stats', 'updated_at'])

        for season, era, war in [
            (2021, '3.30', '1.2'),
            (2020, '0.83', '0.8'),
            (2019, '7.06', '-0.4'),
            (2018, '3.23', '1.4'),
        ]:
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=season,
                team='ATL',
                player_name='A.J. Minter',
                name_ascii='A.J. Minter',
                external_player_id='18655',
                mlbam_id='621345',
                age=28,
                war=float(war),
                raw_stats={
                    'Season': str(season),
                    'Name': 'A.J. Minter',
                    'Team': 'ATL',
                    'Position': 'P',
                    'Height': '6\' 0"',
                    'Weight': '215',
                    'Bats': 'L',
                    'Throws': 'L',
                    'ERA': era,
                    'WAR': war,
                    'PlayerId': '18655',
                    'MLBAMID': '621345',
                },
            )

        response = self.client.get('/api/rosters/ATL/players/621345/')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['roster_season'], 2022)
        self.assertEqual(payload['history_count'], 5)
        self.assertEqual(
            [row['season'] for row in payload['history']],
            [2022, 2021, 2020, 2019, 2018],
        )
        self.assertEqual(payload['player']['photo_url'], '/api/rosters/ATL/players/621345/photo/?season=2022')
        self.assertEqual(payload['player']['pitching']['source_player_id'], '621345')
        self.assertEqual(payload['player']['pitching']['source_external_player_id'], '18655')
        self.assertEqual(payload['history'][0]['pitching']['stats']['games'], 61)

    def test_roster_player_detail_endpoint_includes_similar_players(self):
        stat_line = MLBApiStatLine.objects.get(stat_view='pitching', season=2022, mlbam_id='621345')
        stat_line.external_player_id = '18655'
        stat_line.raw_stats['PlayerId'] = '18655'
        stat_line.save(update_fields=['external_player_id', 'raw_stats', 'updated_at'])

        MLBApiSimilarPlayer.objects.create(
            stat_view='pitching',
            source_player_name='A.J. Minter',
            source_name_ascii='ajminter',
            source_mlbam_id='621345',
            source_external_player_id='18655',
            similar_player_name='Joely Rodríguez',
            similar_name_ascii='joelyrodriguez',
            similar_mlbam_id='570257',
            similar_external_player_id='11487',
            similar_team='NYM',
            similarity_score=51,
            rank=1,
        )
        MLBApiRecommendedSimilarPlayer.objects.create(
            stat_view='pitching',
            source_player_name='A.J. Minter',
            source_name_ascii='ajminter',
            source_mlbam_id='621345',
            source_external_player_id='18655',
            source_player_position='P',
            source_player_age=28,
            similar_player_name='Daniel Norris',
            similar_name_ascii='danielnorris',
            similar_mlbam_id='596057',
            similar_external_player_id='14871',
            similar_team='DET',
            similar_player_position='P',
            similar_player_age=29,
            similarity_score='0.876543210',
            rank=1,
        )

        response = self.client.get('/api/rosters/ATL/players/621345/?season=2022')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['euclidean_similar_players']['batting'], [])
        self.assertEqual(len(payload['euclidean_similar_players']['pitching']), 1)
        self.assertEqual(payload['euclidean_similar_players']['pitching'][0]['player_name'], 'Joely Rodríguez')
        self.assertEqual(payload['euclidean_similar_players']['pitching'][0]['similarity_score'], 51)
        self.assertEqual(payload['tabnet_similar_players']['batting'], [])
        self.assertEqual(len(payload['tabnet_similar_players']['pitching']), 1)
        self.assertEqual(payload['tabnet_similar_players']['pitching'][0]['player_name'], 'Daniel Norris')
        self.assertAlmostEqual(payload['tabnet_similar_players']['pitching'][0]['similarity_score'], 0.87654321)
        self.assertEqual(payload['similar_players'], payload['euclidean_similar_players'])
        self.assertEqual(payload['similar_player_recommendations'], payload['tabnet_similar_players'])
        self.assertIsNone(payload['predicted_aav'])
        self.assertEqual(payload['predicted_aav_meta']['source'], 'M1')
        self.assertFalse(payload['predicted_aav_meta']['available'])

    def test_roster_player_detail_endpoint_includes_predicted_aav_for_matching_batter(self):
        MLBApiAavPrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M1',
            source_file='M1_2022_Predictions.xlsx',
            player_name='Dansby Swanson',
            name_ascii='dansbyswanson',
            team_code_raw='ATL',
            position_raw='SS',
            actual_aav_millions='21.25',
            predicted_aav_millions='19.75',
            prediction_error_millions='-1.50',
        )
        MLBApiPerformanceValuePrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M2',
            source_file='M2_WAR_Value_2022_Full.xlsx',
            player_name='Dansby Swanson',
            name_ascii='dansbyswanson',
            player_type='hitter',
            current_team='ATL',
            target_team='ATL',
            war_2022='6.400',
            predicted_war_avg='3.125',
            actual_war_avg='2.500',
            dollars_per_war_millions='7.10',
            predicted_value_millions='22.19',
        )

        response = self.client.get('/api/rosters/ATL/players/621020/?season=2022')

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['predicted_aav'], 19.75)
        self.assertEqual(payload['predicted_market_value'], 19.75)
        self.assertEqual(payload['predicted_performance_value'], 22.19)
        self.assertEqual(payload['predicted_performance_value_meta']['target_team'], 'ATL')
        self.assertEqual(payload['predicted_performance_value_meta']['basis'], 'war_value_aav')
        self.assertEqual(
            payload['predicted_aav_meta'],
            {
                'source': 'M1',
                'source_file': 'M1_2022_Predictions.xlsx',
                'season': 2022,
                'view': 'batting',
                'unit': 'USD_M',
                'available': True,
            },
        )

    def test_roster_player_detail_endpoint_rejects_external_player_id_lookup(self):
        stat_line = MLBApiStatLine.objects.get(stat_view='pitching', season=2022, mlbam_id='621345')
        stat_line.external_player_id = '18655'
        stat_line.raw_stats['PlayerId'] = '18655'
        stat_line.save(update_fields=['external_player_id', 'raw_stats', 'updated_at'])

        response = self.client.get('/api/rosters/ATL/players/18655/?season=2022')

        self.assertEqual(response.status_code, 404)


class LoadTeamApiStatsCommandTests(TestCase):
    def test_command_loads_csv_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            pitching_csv_body = (
                'Season,Name,Team,HR/9,K%,BB,IP,FIP,GS,G,Age,WAR,NameASCII,PlayerId,MLBAMID\n'
                '2023,Kyle Gibson,BAL,1.078125,0.19454771,55,192.0,4.13,33,33,35,2.66114687919617,Kyle Gibson,10123,502043\n'
            )
            batting_csv_body = (
                'Season,Name,Team,G,PA,HR,ISO,BB%,K%,wOBA,wRC+,SB,BABIP,Age,WAR,NameASCII,PlayerId,MLBAMID\n'
                '2023,Adley Rutschman,BAL,154,687,20,0.192,0.133,0.178,0.374,128,1,0.317,25,5.4,Adley Rutschman,20001,668939\n'
            )
            (data_dir / 'pitching.csv').write_text(pitching_csv_body, encoding='utf-8')
            (data_dir / 'batting.csv').write_text(batting_csv_body, encoding='utf-8')

            call_command('load_team_api_stats', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiStatLine.objects.count(), 2)
        self.assertTrue(MLBApiStatLine.objects.filter(stat_view='pitching', mlbam_id='502043').exists())
        self.assertTrue(MLBApiStatLine.objects.filter(stat_view='batting', mlbam_id='668939').exists())

    def test_command_loads_snake_case_csv_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            pitching_csv_body = (
                'name,season,position,age,height,weight,bats,throws,debut_year,contract_value,war,era,fip,whip,k_per_9,bb_per_9,ip,so,x_era,x_fip,lob_pct,babip,hr_per_9,velocity,player_id\n'
                'A.J. Minter,2022,P,28,6\' 0\",215,L,L,2017,,2.0,2.06,2.53,0.96,10.67,3.09,70.0,83,2.65,2.48,81.2,0.245,0.64,96.4,621345\n'
            )
            batting_csv_body = (
                'name,team,position,age,height,weight,bats,throws,debut_year,contract_value,war,avg,ops,hr,rbi,wrc_plus,woba,babip,iso,bb_pct,k_pct,exit_velocity,launch_angle,season,player_id\n'
                'Dansby Swanson,ATL,SS,28,6\' 1\",190,R,R,2016,,6.4,0.277,0.776,25,96,116,0.330,0.348,0.179,0.087,0.261,89.5,14.2,2022,621020\n'
            )
            legacy_pitching_csv_body = (
                'Season,Name,Team,HR/9,K%,BB,IP,FIP,GS,G,Age,WAR,NameASCII,PlayerId,MLBAMID\n'
                '2022,A.J. Minter,ATL,0.64,0.313,24,70.0,2.53,0,61,28,2.0,A.J. Minter,30001,621345\n'
            )
            legacy_batting_csv_body = (
                'Season,Name,Team,G,PA,HR,ISO,BB%,K%,wOBA,wRC+,SB,BABIP,Age,WAR,NameASCII,PlayerId,MLBAMID\n'
                '2022,Dansby Swanson,ATL,162,696,25,0.179,0.087,0.261,0.330,116,18,0.348,28,6.4,Dansby Swanson,30002,621020\n'
            )
            (data_dir / 'pitchers_2018_2022.csv').write_text(pitching_csv_body, encoding='utf-8')
            (data_dir / 'batters_2018_2022.csv').write_text(batting_csv_body, encoding='utf-8')
            (data_dir / 'pitching.csv').write_text(legacy_pitching_csv_body, encoding='utf-8')
            (data_dir / 'batting.csv').write_text(legacy_batting_csv_body, encoding='utf-8')

            call_command('load_team_api_stats', '--replace', base_dir=str(temp_dir), verbosity=0)

        pitcher = MLBApiStatLine.objects.get(stat_view='pitching', mlbam_id='621345')
        batter = MLBApiStatLine.objects.get(stat_view='batting', mlbam_id='621020')
        self.assertEqual(pitcher.team, 'ATL')
        self.assertEqual(pitcher.external_player_id, '30001')
        self.assertEqual(pitcher.raw_stats['ERA'], '2.06')
        self.assertEqual(pitcher.raw_stats['K/9'], '10.67')
        self.assertEqual(batter.raw_stats['PA'], '696')
        self.assertEqual(batter.raw_stats['AVG'], '0.277')
        self.assertEqual(batter.raw_stats['RBI'], '96')


class LoadApiSimilarPlayersCommandTests(TestCase):
    def test_command_loads_similar_csv_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            (data_dir / 'similar_batters_2018_2022.csv').write_text(
                'name,rank_1_name,rank_1_similarity_score,rank_2_name,rank_2_similarity_score,rank_3_name,rank_3_similarity_score\n'
                'Dansby Swanson,Francisco Lindor,80,Carlos Correa,78,Trea Turner,75\n',
                encoding='utf-8',
            )
            (data_dir / 'similar_pitchers_2018_2022.csv').write_text(
                'name,rank_1_name,rank_1_similarity_score,rank_2_name,rank_2_similarity_score,rank_3_name,rank_3_similarity_score\n'
                'A.J. Minter,Matthew Boyd,48,Joely Rodriguez,46,Daniel Norris,44\n',
                encoding='utf-8',
            )
            (data_dir / 'batters_recommendations.csv').write_text(
                'query_player_id,query_name,query_position,query_age,rec_1_player_id,rec_1_name,rec_1_position,rec_1_age,rec_1_similarity,rec_2_player_id,rec_2_name,rec_2_position,rec_2_age,rec_2_similarity,rec_3_player_id,rec_3_name,rec_3_position,rec_3_age,rec_3_similarity\n'
                '621020,Dansby Swanson,SS,28,596019,Francisco Lindor,SS,28,0.812345678,622491,Carlos Correa,SS,28,0.801234567,607208,Trea Turner,SS,29,0.792345678\n',
                encoding='utf-8',
            )
            (data_dir / 'pitchers_recommendations.csv').write_text(
                'query_player_id,query_name,query_position,query_age,rec_1_player_id,rec_1_name,rec_1_position,rec_1_age,rec_1_similarity,rec_2_player_id,rec_2_name,rec_2_position,rec_2_age,rec_2_similarity,rec_3_player_id,rec_3_name,rec_3_position,rec_3_age,rec_3_similarity\n'
                '621345,A.J. Minter,P,28,571510,Matthew Boyd,P,31,0.956789123,570257,Joely Rodriguez,P,30,0.934567891,596057,Daniel Norris,P,29,0.912345678\n',
                encoding='utf-8',
            )

            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=2022,
                team='ATL',
                player_name='A.J. Minter',
                name_ascii='A.J. Minter',
                external_player_id='18655',
                mlbam_id='621345',
                age=28,
                war=2.0,
                raw_stats={'PlayerId': '18655', 'MLBAMID': '621345'},
            )
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=2022,
                team='DET',
                player_name='Matthew Boyd',
                name_ascii='Matthew Boyd',
                external_player_id='15440',
                mlbam_id='571510',
                age=31,
                war=1.4,
                raw_stats={'PlayerId': '15440', 'MLBAMID': '571510'},
            )
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=2022,
                team='NYM',
                player_name='Joely Rodriguez',
                name_ascii='Joely Rodriguez',
                external_player_id='11487',
                mlbam_id='570257',
                age=30,
                war=0.8,
                raw_stats={'PlayerId': '11487', 'MLBAMID': '570257'},
            )
            MLBApiStatLine.objects.create(
                stat_view='pitching',
                season=2022,
                team='DET',
                player_name='Daniel Norris',
                name_ascii='Daniel Norris',
                external_player_id='14871',
                mlbam_id='596057',
                age=29,
                war=0.5,
                raw_stats={'PlayerId': '14871', 'MLBAMID': '596057'},
            )
            MLBApiStatLine.objects.create(
                stat_view='batting',
                season=2022,
                team='ATL',
                player_name='Dansby Swanson',
                name_ascii='Dansby Swanson',
                external_player_id='621020',
                mlbam_id='621020',
                age=28,
                war=6.4,
                raw_stats={'PlayerId': '621020', 'MLBAMID': '621020'},
            )
            MLBApiStatLine.objects.create(
                stat_view='batting',
                season=2022,
                team='NYM',
                player_name='Francisco Lindor',
                name_ascii='Francisco Lindor',
                external_player_id='12916',
                mlbam_id='596019',
                age=28,
                war=6.8,
                raw_stats={'PlayerId': '12916', 'MLBAMID': '596019'},
            )
            MLBApiStatLine.objects.create(
                stat_view='batting',
                season=2022,
                team='MIN',
                player_name='Carlos Correa',
                name_ascii='Carlos Correa',
                external_player_id='14162',
                mlbam_id='622491',
                age=28,
                war=4.4,
                raw_stats={'PlayerId': '14162', 'MLBAMID': '622491'},
            )
            MLBApiStatLine.objects.create(
                stat_view='batting',
                season=2022,
                team='LAD',
                player_name='Trea Turner',
                name_ascii='Trea Turner',
                external_player_id='16252',
                mlbam_id='607208',
                age=29,
                war=6.3,
                raw_stats={'PlayerId': '16252', 'MLBAMID': '607208'},
            )

            call_command('load_api_similar_players', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiSimilarPlayer.objects.count(), 6)
        self.assertEqual(MLBApiRecommendedSimilarPlayer.objects.count(), 6)
        pitcher_similar = MLBApiSimilarPlayer.objects.get(stat_view='pitching', source_name_ascii='ajminter', rank=1)
        batter_similar = MLBApiSimilarPlayer.objects.get(stat_view='batting', source_name_ascii='dansbyswanson', rank=1)
        pitcher_recommendation = MLBApiRecommendedSimilarPlayer.objects.get(
            stat_view='pitching',
            source_name_ascii='ajminter',
            rank=1,
        )
        batter_recommendation = MLBApiRecommendedSimilarPlayer.objects.get(
            stat_view='batting',
            source_name_ascii='dansbyswanson',
            rank=1,
        )
        self.assertEqual(pitcher_similar.source_mlbam_id, '621345')
        self.assertEqual(pitcher_similar.similar_external_player_id, '15440')
        self.assertEqual(batter_similar.similar_mlbam_id, '596019')
        self.assertEqual(batter_similar.similarity_score, 80)
        self.assertEqual(pitcher_recommendation.similar_external_player_id, '15440')
        self.assertEqual(pitcher_recommendation.similar_player_position, 'P')
        self.assertEqual(pitcher_recommendation.similar_player_age, 31)
        self.assertEqual(str(pitcher_recommendation.similarity_score), '0.956789123')
        self.assertEqual(batter_recommendation.similar_mlbam_id, '596019')
        self.assertEqual(str(batter_recommendation.similarity_score), '0.812345678')


class LoadApiAavPredictionsCommandTests(TestCase):
    def test_command_loads_workbook_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '2022_M1_Predictions'
            sheet.append(('2022년 FA 타자 — M1 시장 기반 모델 예측값', None, None, None, None, None, None))
            sheet.append(('Ridge Regression | 학습: 2002~2021 | 예측: 2022년 FA 타자 | 단위: $M (백만달러)', None, None, None, None, None, None))
            sheet.append(('연도', '선수명', '포지션', '계약팀', '실제 AAV ($M)', 'M1 예측값 ($M)', '오차 ($M)'))
            sheet.append((2022, 'Carlos Correa', 'SS', 'MIN', 35.1, 24.08, -11.02))
            sheet.append((2022, 'Corey Seager', 'SS', 'TEX', 32.5, 16.34, -16.16))
            sheet.append((2021, 'Ignore Me', 'SS', 'NYY', 20.0, 18.0, -2.0))
            sheet.append(('총 선수 수', '2', None, None, None, None, None))
            workbook.save(data_dir / 'M1_2022_Predictions.xlsx')

            call_command('load_api_aav_predictions', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiAavPrediction.objects.count(), 2)
        correa = MLBApiAavPrediction.objects.get(name_ascii='carloscorrea')
        seager = MLBApiAavPrediction.objects.get(name_ascii='coreyseager')
        self.assertEqual(correa.player_name, 'Carlos Correa')
        self.assertEqual(correa.team_code_raw, 'MIN')
        self.assertEqual(float(correa.predicted_aav_millions), 24.08)
        self.assertEqual(float(correa.prediction_error_millions), -11.02)
        self.assertEqual(seager.position_raw, 'SS')

    def test_command_replace_reloads_rows_without_duplicates(self):
        MLBApiAavPrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M1',
            source_file='M1_2022_Predictions.xlsx',
            player_name='Carlos Correa',
            name_ascii='carloscorrea',
            team_code_raw='MIN',
            position_raw='SS',
            actual_aav_millions='35.10',
            predicted_aav_millions='20.00',
            prediction_error_millions='-15.10',
        )

        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '2022_M1_Predictions'
            sheet.append(('2022년 FA 타자 — M1 시장 기반 모델 예측값', None, None, None, None, None, None))
            sheet.append(('Ridge Regression | 학습: 2002~2021 | 예측: 2022년 FA 타자 | 단위: $M (백만달러)', None, None, None, None, None, None))
            sheet.append(('연도', '선수명', '포지션', '계약팀', '실제 AAV ($M)', 'M1 예측값 ($M)', '오차 ($M)'))
            sheet.append((2022, 'Carlos Correa', 'SS', 'MIN', 35.1, 24.08, -11.02))
            workbook.save(data_dir / 'M1_2022_Predictions.xlsx')

            call_command('load_api_aav_predictions', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiAavPrediction.objects.count(), 1)
        correa = MLBApiAavPrediction.objects.get(name_ascii='carloscorrea')
        self.assertEqual(float(correa.predicted_aav_millions), 24.08)


class LoadApiPerformanceValuesCommandTests(TestCase):
    def test_command_loads_workbook_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'player_team_values'
            sheet.append(('2022 FA player team values', None, None, None, None, None, None, None, None, None, None))
            sheet.append((
                'player_name_key',
                'Name',
                'player_type',
                'Season',
                'current_team',
                'target_team',
                'WAR(2022)',
                'pred_WAR(23~25avg)',
                'actual_WAR(23~25avg)',
                '$/WAR($M)',
                'M2 value($M)',
            ))
            sheet.append(('dansbyswanson', 'Dansby Swanson', 'hitter', 2022, 'ATL', 'ATL', 6.4, 3.125, 2.5, 7.1, 22.19))
            sheet.append(('dansbyswanson', 'Dansby Swanson', 'hitter', 2022, 'ATL', 'TEX', 6.4, 3.125, 2.5, 9.56, 29.88))
            sheet.append(('aaronnola', 'Aaron Nola', 'pitcher', 2022, 'PHI', 'ATL', 6.296, 3.185, 2.624, 7.1, 22.61))
            sheet.append(('ignoreme', 'Ignore Me', 'hitter', 2021, 'NYY', 'ATL', 1.0, 1.0, 1.0, 7.1, 7.1))
            workbook.save(data_dir / 'M2_WAR_Value_2022_Full.xlsx')

            call_command('load_api_performance_values', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiPerformanceValuePrediction.objects.count(), 3)
        swanson_atl = MLBApiPerformanceValuePrediction.objects.get(name_ascii='dansbyswanson', target_team='ATL')
        swanson_tex = MLBApiPerformanceValuePrediction.objects.get(name_ascii='dansbyswanson', target_team='TEX')
        nola = MLBApiPerformanceValuePrediction.objects.get(name_ascii='aaronnola')
        self.assertEqual(swanson_atl.stat_view, 'batting')
        self.assertEqual(nola.stat_view, 'pitching')
        self.assertEqual(float(swanson_atl.predicted_value_millions), 22.19)
        self.assertEqual(float(swanson_tex.dollars_per_war_millions), 9.56)

    def test_command_replace_reloads_rows_without_duplicates(self):
        MLBApiPerformanceValuePrediction.objects.create(
            stat_view='batting',
            season=2022,
            source_label='M2',
            source_file='M2_WAR_Value_2022_Full.xlsx',
            player_name='Dansby Swanson',
            name_ascii='dansbyswanson',
            player_type='hitter',
            current_team='ATL',
            target_team='ATL',
            predicted_value_millions='20.00',
        )

        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data'
            data_dir.mkdir(parents=True, exist_ok=True)

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'player_team_values'
            sheet.append(('2022 FA player team values', None, None, None, None, None, None, None, None, None, None))
            sheet.append((
                'player_name_key',
                'Name',
                'player_type',
                'Season',
                'current_team',
                'target_team',
                'WAR(2022)',
                'pred_WAR(23~25avg)',
                'actual_WAR(23~25avg)',
                '$/WAR($M)',
                'M2 value($M)',
            ))
            sheet.append(('dansbyswanson', 'Dansby Swanson', 'hitter', 2022, 'ATL', 'ATL', 6.4, 3.125, 2.5, 7.1, 22.19))
            workbook.save(data_dir / 'M2_WAR_Value_2022_Full.xlsx')

            call_command('load_api_performance_values', '--replace', base_dir=str(temp_dir), verbosity=0)

        self.assertEqual(MLBApiPerformanceValuePrediction.objects.count(), 1)
        swanson = MLBApiPerformanceValuePrediction.objects.get(name_ascii='dansbyswanson')
        self.assertEqual(float(swanson.predicted_value_millions), 22.19)


class LoadRosterPhotosCommandTests(TestCase):
    def test_command_loads_photo_rows_into_db(self):
        with TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / 'data' / 'Atlanta Braves'
            nested_dir = data_dir / 'Atlanta Braves'
            nested_dir.mkdir(parents=True, exist_ok=True)
            (nested_dir / 'A.J. Minter.jpeg').write_bytes(b'photo-bytes')

            call_command('load_roster_photos', '--replace', base_dir=str(temp_dir), verbosity=0)

        photo = MLBRosterPhoto.objects.get(team_name='Atlanta Braves', normalized_player_name='ajminter')
        self.assertEqual(photo.player_name, 'A.J. Minter')
        self.assertEqual(photo.content_type, 'image/jpeg')
        self.assertEqual(bytes(photo.image_data), b'photo-bytes')
