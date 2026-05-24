from decimal import Decimal
from pathlib import Path
import unicodedata

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from mlb.models import MLBApiPerformanceValuePrediction, MLBApiStatLine


DEFAULT_SOURCE_LABEL = 'M2'
DEFAULT_SOURCE_FILE = 'M2_WAR_Value_2022_Full.xlsx'
DEFAULT_SHEET_NAME = 'player_team_values'
DEFAULT_SEASON = 2022


def _normalize_name(value):
    normalized = unicodedata.normalize('NFKD', value or '')
    ascii_only = normalized.encode('ascii', 'ignore').decode('ascii')
    return ''.join(ch.lower() for ch in ascii_only if ch.isalnum())


def _to_decimal(value, places):
    if value in (None, ''):
        return None
    quantum = Decimal('1').scaleb(-places)
    return Decimal(str(value)).quantize(quantum)


def _stat_view_for_player_type(player_type):
    normalized = str(player_type or '').strip().lower()
    if normalized == 'hitter':
        return MLBApiStatLine.VIEW_BATTING
    if normalized == 'pitcher':
        return MLBApiStatLine.VIEW_PITCHING
    return None


class Command(BaseCommand):
    help = 'Load API-facing WAR-based performance values from the M2 workbook into a dedicated table.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--replace',
            action='store_true',
            help='Delete existing rows for the same source label before loading.',
        )
        parser.add_argument(
            '--base-dir',
            help='Optional base directory containing the data/ folder.',
        )

    def handle(self, *args, **options):
        base_dir = Path(options['base_dir']) if options.get('base_dir') else Path(settings.BASE_DIR)
        workbook_path = base_dir / 'data' / DEFAULT_SOURCE_FILE
        if not workbook_path.exists():
            raise CommandError(f'Performance value workbook not found: {workbook_path}')

        if options['replace']:
            deleted_count, _ = MLBApiPerformanceValuePrediction.objects.filter(
                source_label=DEFAULT_SOURCE_LABEL
            ).delete()
            self.stdout.write(f'Deleted {deleted_count} existing {DEFAULT_SOURCE_LABEL} rows.')

        prediction_objects = self._load_workbook(workbook_path)

        with transaction.atomic():
            if prediction_objects:
                MLBApiPerformanceValuePrediction.objects.bulk_create(
                    prediction_objects,
                    batch_size=500,
                    update_conflicts=True,
                    unique_fields=['season', 'stat_view', 'name_ascii', 'target_team'],
                    update_fields=[
                        'source_label',
                        'source_file',
                        'player_name',
                        'player_type',
                        'current_team',
                        'war_2022',
                        'predicted_war_avg',
                        'actual_war_avg',
                        'dollars_per_war_millions',
                        'predicted_value_millions',
                        'updated_at',
                    ],
                )

        self.stdout.write(
            self.style.SUCCESS(
                f'Loaded or updated {len(prediction_objects)} {DEFAULT_SOURCE_LABEL} performance value rows.'
            )
        )

    def _load_workbook(self, workbook_path):
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if DEFAULT_SHEET_NAME not in workbook.sheetnames:
            workbook.close()
            raise CommandError(f'Worksheet not found in M2 workbook: {DEFAULT_SHEET_NAME}')

        worksheet = workbook[DEFAULT_SHEET_NAME]
        objects = []
        for row in worksheet.iter_rows(min_row=3, values_only=True):
            season = row[3] if len(row) > 3 else None
            player_name = (row[1] if len(row) > 1 else '') or ''
            target_team = ((row[5] if len(row) > 5 else '') or '').strip().upper()
            stat_view = _stat_view_for_player_type(row[2] if len(row) > 2 else None)
            if season != DEFAULT_SEASON or not str(player_name).strip() or not target_team or stat_view is None:
                continue

            objects.append(
                MLBApiPerformanceValuePrediction(
                    stat_view=stat_view,
                    season=DEFAULT_SEASON,
                    source_label=DEFAULT_SOURCE_LABEL,
                    source_file=DEFAULT_SOURCE_FILE,
                    player_name=str(player_name).strip(),
                    name_ascii=_normalize_name(str(player_name).strip()),
                    player_type=str((row[2] if len(row) > 2 else '') or '').strip(),
                    current_team=((row[4] if len(row) > 4 else '') or '').strip().upper(),
                    target_team=target_team,
                    war_2022=_to_decimal(row[6] if len(row) > 6 else None, 3),
                    predicted_war_avg=_to_decimal(row[7] if len(row) > 7 else None, 3),
                    actual_war_avg=_to_decimal(row[8] if len(row) > 8 else None, 3),
                    dollars_per_war_millions=_to_decimal(row[9] if len(row) > 9 else None, 2),
                    predicted_value_millions=_to_decimal(row[10] if len(row) > 10 else None, 2),
                )
            )
        workbook.close()
        return objects
