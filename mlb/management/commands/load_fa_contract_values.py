import unicodedata
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from mlb.models import MLBApiStatLine

DEFAULT_SOURCE_FILE = 'MLB-Free Agency 1991-2026_consolidated.xlsx'


def _normalize_xlsx_name(value):
    """xlsx Player 이름을 MLBApiStatLine.name_ascii 형식으로 변환.

    load_mock_team_api_data._name_ascii()와 동일한 NFKD+ASCII 변환 후
    xlsx에서 올 수 있는 하이픈을 공백으로 치환한다.
    """
    text = str(value).strip() if value is not None else ''
    normalized = unicodedata.normalize('NFKD', text)
    ascii_str = normalized.encode('ascii', 'ignore').decode('ascii')
    return ascii_str.replace('-', ' ')


class Command(BaseCommand):
    help = (
        'MLB-Free Agency 1991-2026_consolidated.xlsx에서 선수별 최신 signed 계약 AAV를 읽어 '
        'MLBApiStatLine.raw_stats["ContractValue"]에 적재한다. '
        '실행 전 load_mock_team_api_data로 MLBApiStatLine 행이 존재해야 한다.'
    )

    def add_arguments(self, parser):
        parser.add_argument(
            '--replace',
            action='store_true',
            help='기존 ContractValue를 모든 행에서 제거한 뒤 재적재.',
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='실제 저장 없이 매칭 결과만 출력.',
        )
        parser.add_argument(
            '--base-dir',
            help='data/ 폴더를 포함하는 프로젝트 루트 경로 (기본: settings.BASE_DIR).',
        )

    def handle(self, *args, **options):
        base_dir = Path(options['base_dir']) if options.get('base_dir') else Path(settings.BASE_DIR)
        workbook_path = base_dir / 'data' / DEFAULT_SOURCE_FILE
        if not workbook_path.exists():
            raise CommandError(f'FA contract workbook not found: {workbook_path}')

        if options['replace']:
            self._clear_contract_values()

        lookup = self._load_fa_lookup(workbook_path)
        self.stdout.write(f'Loaded {len(lookup)} distinct signed player contracts from xlsx.')

        updated = self._apply_contract_values(lookup, dry_run=options['dry_run'])

        if not options['dry_run']:
            self.stdout.write(self.style.SUCCESS(f'Done. Updated {updated} stat line rows.'))

    def _clear_contract_values(self):
        lines_with_cv = [
            line for line in MLBApiStatLine.objects.all()
            if 'ContractValue' in (line.raw_stats or {})
        ]
        for line in lines_with_cv:
            raw = dict(line.raw_stats)
            raw.pop('ContractValue', None)
            line.raw_stats = raw
        with transaction.atomic():
            MLBApiStatLine.objects.bulk_update(lines_with_cv, ['raw_stats'], batch_size=500)
        self.stdout.write(f'Cleared ContractValue from {len(lines_with_cv)} rows.')

    def _load_fa_lookup(self, workbook_path):
        """xlsx에서 선수별 최신 signed 계약 AAV를 추출한다.

        Returns: {name_ascii: aav_millions (float)}
        """
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        best = {}  # name_ascii -> (year, aav_millions)
        try:
            ws = wb.active
            for row in ws.iter_rows(min_row=2, max_col=14, values_only=True):
                (year, status, player, pos, age, qo, old_club, new_club,
                 years, guarantee, term, option, opt_out, aav) = row

                if str(status or '').strip().lower() != 'signed':
                    continue
                if not player or aav is None:
                    continue

                try:
                    aav_float = float(aav)
                    year_int = int(year) if year is not None else 0
                except (TypeError, ValueError):
                    continue

                aav_millions = round(aav_float / 1_000_000, 4)
                na = _normalize_xlsx_name(str(player).strip())
                if not na:
                    continue

                existing = best.get(na)
                if existing is None or year_int > existing[0]:
                    best[na] = (year_int, aav_millions)
        finally:
            wb.close()

        return {na: v[1] for na, v in best.items()}

    def _apply_contract_values(self, lookup, dry_run):
        all_lines = list(MLBApiStatLine.objects.all())
        to_update = []
        matched_players = set()

        for line in all_lines:
            na = line.name_ascii or ''
            if na not in lookup:
                continue
            raw = dict(line.raw_stats or {})
            raw['ContractValue'] = str(lookup[na])
            line.raw_stats = raw
            to_update.append(line)
            matched_players.add(na)

        skipped = len(all_lines) - len(to_update)

        if dry_run:
            self.stdout.write(
                f'[dry-run] Would update {len(to_update)} rows '
                f'({len(matched_players)} players). '
                f'No contract data for {skipped} rows.'
            )
            for na in sorted(matched_players)[:30]:
                self.stdout.write(f'  {na!r} -> {lookup[na]:.4f}M')
            return len(to_update)

        with transaction.atomic():
            MLBApiStatLine.objects.bulk_update(to_update, ['raw_stats'], batch_size=500)

        self.stdout.write(
            f'Updated {len(to_update)} rows ({len(matched_players)} players). '
            f'No contract data: {skipped} rows skipped.'
        )
        return len(to_update)
