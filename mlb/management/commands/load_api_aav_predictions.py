from decimal import Decimal
from pathlib import Path
import unicodedata

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from mlb.models import MLBApiAavPrediction, MLBApiStatLine


DEFAULT_SOURCE_LABEL = 'M1'
DEFAULT_SOURCE_FILE = 'M1_2022_Predictions.xlsx'
DEFAULT_SEASON = 2022


def _normalize_name(value):
    normalized = unicodedata.normalize('NFKD', value or '')
    ascii_only = normalized.encode('ascii', 'ignore').decode('ascii')
    return ''.join(ch.lower() for ch in ascii_only if ch.isalnum())


def _to_decimal(value):
    if value in (None, ''):
        return None
    return Decimal(str(value)).quantize(Decimal('0.01'))


class Command(BaseCommand):
    help = 'Load API-facing AAV predictions from the M1 workbook into a dedicated table.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--replace',
            action='store_true',
            help='Delete existing rows for the same source label before loading.',
        )
        parser.add_argument(
            '--base-dir',
            help='Optional base directory containing the data/ folder.',
        )

    def handle(self, *args, **options):
        base_dir = Path(options['base_dir']) if options.get('base_dir') else Path(settings.BASE_DIR)
        workbook_path = base_dir / 'data' / DEFAULT_SOURCE_FILE
        if not workbook_path.exists():
            raise CommandError(f'AAV prediction workbook not found: {workbook_path}')

        if options['replace']:
            deleted_count, _ = MLBApiAavPrediction.objects.filter(source_label=DEFAULT_SOURCE_LABEL).delete()
            self.stdout.write(f'Deleted {deleted_count} existing {DEFAULT_SOURCE_LABEL} rows.')

        prediction_objects = self._load_workbook(workbook_path)

        with transaction.atomic():
            if prediction_objects:
                MLBApiAavPrediction.objects.bulk_create(
                    prediction_objects,
                    batch_size=500,
                    update_conflicts=True,
                    unique_fields=['season', 'stat_view', 'name_ascii'],
                    update_fields=[
                        'source_label',
                        'source_file',
                        'player_name',
                        'team_code_raw',
                        'position_raw',
                        'actual_aav_millions',
                        'predicted_aav_millions',
                        'prediction_error_millions',
                        'updated_at',
                    ],
                )

        self.stdout.write(
            self.style.SUCCESS(
                f'Loaded or updated {len(prediction_objects)} {DEFAULT_SOURCE_LABEL} AAV prediction rows.'
            )
        )

    def _load_workbook(self, workbook_path):
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        objects = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(min_row=4, values_only=True):
                season = row[0] if len(row) > 0 else None
                player_name = (row[1] if len(row) > 1 else '') or ''
                if season != DEFAULT_SEASON or not str(player_name).strip():
                    continue

                objects.append(
                    MLBApiAavPrediction(
                        stat_view=MLBApiStatLine.VIEW_BATTING,
                        season=DEFAULT_SEASON,
                        source_label=DEFAULT_SOURCE_LABEL,
                        source_file=DEFAULT_SOURCE_FILE,
                        player_name=str(player_name).strip(),
                        name_ascii=_normalize_name(str(player_name).strip()),
                        position_raw=((row[2] if len(row) > 2 else '') or '').strip(),
                        team_code_raw=((row[3] if len(row) > 3 else '') or '').strip(),
                        actual_aav_millions=_to_decimal(row[4] if len(row) > 4 else None),
                        predicted_aav_millions=_to_decimal(row[5] if len(row) > 5 else None),
                        prediction_error_millions=_to_decimal(row[6] if len(row) > 6 else None),
                    )
                )
        workbook.close()
        return objects
