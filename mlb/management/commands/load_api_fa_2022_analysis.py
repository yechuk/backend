from decimal import Decimal
from pathlib import Path
import re
import unicodedata

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from mlb.models import MLBApiFa2022Analysis, MLBApiStatLine


DEFAULT_SOURCE_FILE = 'FA_2022_최종분석_v3.xlsx'
DEFAULT_SEASON = 2022
SHEET_VIEW_MAP = {
    '타자_메인': MLBApiStatLine.VIEW_BATTING,
    '투수_메인': MLBApiStatLine.VIEW_PITCHING,
}
M2_SHEET_MAP = {
    '타자_M2_30팀': MLBApiStatLine.VIEW_BATTING,
    '투수_M2_30팀': MLBApiStatLine.VIEW_PITCHING,
}
SOURCE_HEADERS = [
    '선수명',
    '이전팀',
    '계약팀',
    '포지션',
    '재계약',
    '나이',
    'pred_WAR',
    '실제AAV\n($M)',
    'M1 예측\n($M)',
    'M2\n계약팀',
    'M2\n최대값',
    'M2\n최대팀',
    'M2\n최소값',
    'M2\n최소팀',
    '포지션 희소성',
    'Boras 여부',
    '나이 신호',
    '팀 특수성',
]


def _normalize_name(value):
    normalized = unicodedata.normalize('NFKD', value or '')
    ascii_only = normalized.encode('ascii', 'ignore').decode('ascii')
    return ''.join(ch.lower() for ch in ascii_only if ch.isalnum())


def _to_decimal(value, places):
    if value in (None, ''):
        return None
    quantum = Decimal('1').scaleb(-places)
    return Decimal(str(value)).quantize(quantum)


def _clean_text(value):
    if value is None:
        return ''
    return str(value).strip()


def _parse_is_re_signing(value):
    cleaned = _clean_text(value)
    if cleaned == '재계약':
        return True
    if cleaned == '이적':
        return False
    raise CommandError(f'Unexpected 재계약 value: {value!r}')


def _parse_is_boras(value):
    cleaned = _clean_text(value)
    if 'Non-Boras' in cleaned:
        return False
    if 'Boras' in cleaned:
        return True
    raise CommandError(f'Unexpected Boras 여부 value: {value!r}')


def _parse_position_scarcity(value):
    cleaned = _clean_text(value)
    if '과열 경고' in cleaned:
        level = 'overheated'
    elif '주의' in cleaned:
        level = 'caution'
    elif '안정' in cleaned:
        level = 'stable'
    else:
        raise CommandError(f'Unexpected 포지션 희소성 level: {value!r}')

    match = re.search(r'predWAR\s*([0-9.]+)\+\s*([0-9]+)명', cleaned)
    if match is None:
        raise CommandError(f'Unexpected 포지션 희소성 format: {value!r}')

    return {
        'position_scarcity_level': level,
        'position_scarcity_war_threshold': _to_decimal(match.group(1), 1),
        'position_scarcity_comp_count': int(match.group(2)),
    }


def _parse_age_signal(value):
    cleaned = _clean_text(value)
    if '전성기 프리미엄' in cleaned:
        level = 'prime_premium'
    elif '에이징 리스크' in cleaned:
        level = 'aging_risk'
    elif '해당없음' in cleaned:
        level = 'none'
    else:
        raise CommandError(f'Unexpected 나이 신호 level: {value!r}')

    return {
        'age_signal_level': level,
        'age_signal_is_aging_risk': level == 'aging_risk',
    }


def _parse_team_specificity(value):
    cleaned = _clean_text(value)
    if '데이터없음' in cleaned:
        return {'team_specificity_level': 'no_data', 'team_specificity_dollar_per_war': None}
    elif '해당없음' in cleaned:
        level = 'none'
    elif 'M2 참고 권장' in cleaned:
        level = 'm2_recommended'
    elif 'M2 억제 주의' in cleaned:
        level = 'm2_suppressed'
    else:
        raise CommandError(f'Unexpected 팀 특수성 level: {value!r}')

    match = re.search(r'\$/WAR=\$([0-9.]+)M', cleaned)
    if match is None:
        raise CommandError(f'Unexpected 팀 특수성 format: {value!r}')

    return {
        'team_specificity_level': level,
        'team_specificity_dollar_per_war': _to_decimal(match.group(1), 2),
    }


class Command(BaseCommand):
    help = 'Load API-facing 2022 FA analysis rows from the v3 final workbook.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--replace',
            action='store_true',
            help='Delete existing FA 2022 analysis rows before loading.',
        )
        parser.add_argument(
            '--base-dir',
            help='Optional base directory containing the data/ folder.',
        )

    def handle(self, *args, **options):
        base_dir = Path(options['base_dir']) if options.get('base_dir') else Path(settings.BASE_DIR)
        workbook_path = base_dir / 'data' / DEFAULT_SOURCE_FILE
        if not workbook_path.exists():
            raise CommandError(f'FA 2022 analysis workbook not found: {workbook_path}')

        if options['replace']:
            deleted_count, _ = MLBApiFa2022Analysis.objects.filter(season=DEFAULT_SEASON).delete()
            self.stdout.write(f'Deleted {deleted_count} existing FA 2022 analysis rows.')

        objects = self._load_workbook(workbook_path)

        with transaction.atomic():
            if objects:
                MLBApiFa2022Analysis.objects.bulk_create(
                    objects,
                    batch_size=500,
                    update_conflicts=True,
                    unique_fields=['season', 'stat_view', 'name_ascii', 'previous_team'],
                    update_fields=[
                        'player_name',
                        'contract_team',
                        'position',
                        'is_re_signing',
                        'predicted_war',
                        'actual_aav_millions',
                        'predicted_aav_millions',
                        'position_scarcity_level',
                        'position_scarcity_war_threshold',
                        'position_scarcity_comp_count',
                        'is_boras',
                        'age',
                        'age_signal_level',
                        'age_signal_is_aging_risk',
                        'm2_contract_team_millions',
                        'm2_max_millions',
                        'm2_max_team',
                        'm2_min_millions',
                        'm2_min_team',
                        'team_specificity_level',
                        'team_specificity_dollar_per_war',
                        'm2_by_team',
                        'updated_at',
                    ],
                )

        self.stdout.write(
            self.style.SUCCESS(f'Loaded or updated {len(objects)} FA 2022 analysis rows.')
        )

    def _load_workbook(self, workbook_path):
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        objects = []
        try:
            m2_lookup = self._load_m2_by_team(workbook)

            for sheet_name, stat_view in SHEET_VIEW_MAP.items():
                if sheet_name not in workbook.sheetnames:
                    raise CommandError(f'Worksheet not found in FA 2022 workbook: {sheet_name}')
                worksheet = workbook[sheet_name]
                header_row = next(
                    worksheet.iter_rows(min_row=3, max_row=3, min_col=1, max_col=18, values_only=True)
                )
                cleaned_headers = [_clean_text(value) for value in header_row]
                if cleaned_headers != SOURCE_HEADERS:
                    raise CommandError(
                        f'Unexpected headers in sheet {sheet_name}. '
                        f'Expected {SOURCE_HEADERS!r}, got {cleaned_headers!r}'
                    )

                for row in worksheet.iter_rows(min_row=4, min_col=1, max_col=18, values_only=True):
                    values = list(row)
                    if all(_clean_text(value) == '' for value in values):
                        break

                    player_name = _clean_text(values[0])
                    previous_team = _clean_text(values[1]).upper()
                    contract_team = _clean_text(values[2]).upper()
                    if not player_name or not previous_team:
                        raise CommandError(
                            f'Missing required player_name/previous_team in sheet {sheet_name}: {values!r}'
                        )

                    name_ascii = _normalize_name(player_name)
                    scarcity = _parse_position_scarcity(values[14])
                    age_signal = _parse_age_signal(values[16])
                    team_specificity = _parse_team_specificity(values[17])
                    m2_by_team = m2_lookup.get((stat_view, name_ascii, previous_team))

                    objects.append(
                        MLBApiFa2022Analysis(
                            season=DEFAULT_SEASON,
                            stat_view=stat_view,
                            player_name=player_name,
                            name_ascii=name_ascii,
                            previous_team=previous_team,
                            contract_team=contract_team,
                            position=_clean_text(values[3]),
                            is_re_signing=_parse_is_re_signing(values[4]),
                            age=_to_decimal(values[5], 1),
                            predicted_war=_to_decimal(values[6], 3),
                            actual_aav_millions=_to_decimal(values[7], 2),
                            predicted_aav_millions=_to_decimal(values[8], 2),
                            m2_contract_team_millions=_to_decimal(values[9], 2),
                            m2_max_millions=_to_decimal(values[10], 2),
                            m2_max_team=_clean_text(values[11]).upper(),
                            m2_min_millions=_to_decimal(values[12], 2),
                            m2_min_team=_clean_text(values[13]).upper(),
                            is_boras=_parse_is_boras(values[15]),
                            m2_by_team=m2_by_team,
                            **scarcity,
                            **age_signal,
                            **team_specificity,
                        )
                    )
        finally:
            workbook.close()
        return objects

    def _load_m2_by_team(self, workbook):
        """M2_30팀 시트에서 선수별 30팀 M2 예측값을 읽어 lookup dict로 반환.

        Returns: {(stat_view, name_ascii, previous_team): {"ARI": 21.26, ...}}
        """
        lookup = {}
        for sheet_name, stat_view in M2_SHEET_MAP.items():
            if sheet_name not in workbook.sheetnames:
                self.stdout.write(self.style.WARNING(f'M2 30팀 시트 없음: {sheet_name}'))
                continue
            worksheet = workbook[sheet_name]

            # row 3: 헤더 (선수명, 이전팀, 계약팀, 포지션, pred_WAR, ARI, ATL, ...)
            header_row = list(
                worksheet.iter_rows(min_row=3, max_row=3, min_col=1, max_col=35, values_only=True)
            )[0]
            team_codes = [_clean_text(h) for h in header_row[5:]]  # col 6-35

            for row in worksheet.iter_rows(min_row=4, min_col=1, max_col=35, values_only=True):
                values = list(row)
                player_name = _clean_text(values[0])
                if not player_name or player_name == '팀별 $/WAR ($M)':
                    continue

                previous_team = _clean_text(values[1]).upper()
                name_ascii = _normalize_name(player_name)
                m2_values = {
                    team_codes[i]: round(float(values[5 + i]), 2)
                    for i in range(len(team_codes))
                    if values[5 + i] is not None
                }
                lookup[(stat_view, name_ascii, previous_team)] = m2_values

        return lookup
